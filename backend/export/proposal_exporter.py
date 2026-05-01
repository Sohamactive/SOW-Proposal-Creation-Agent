import os
from typing import Any

from docx import Document as WordDocument
from pylatex import Command, Document as LatexDocument, Itemize, Section, Subsection
from pylatex.errors import CompilerError
from pylatex.utils import escape_latex
from pptx import Presentation

from backend.config import settings


class ProposalExporter:

    def __init__(self):
        os.makedirs(settings.EXPORT_FOLDER, exist_ok=True)

    def export_docx(self, proposal: dict[str, Any], filename: str = "proposal.docx") -> str:

        os.makedirs(settings.EXPORT_FOLDER, exist_ok=True)
        path = os.path.join(settings.EXPORT_FOLDER, filename)

        doc = WordDocument()
        doc.add_heading("Project Proposal", level=0)

        for section, content in proposal.items():
            doc.add_heading(self._format_label(section), level=1)
            self._write_docx_content(doc, content)

        doc.save(path)
        return path

    def _write_docx_content(self, doc: WordDocument, content: Any) -> None:

        if isinstance(content, str):
            for paragraph in self._split_paragraphs(content):
                doc.add_paragraph(paragraph)
            return

        if isinstance(content, dict):
            for key, value in content.items():
                doc.add_heading(self._format_label(key), level=2)
                self._write_docx_content(doc, value)
            return

        if isinstance(content, list):
            for item in content:
                if isinstance(item, dict):
                    for key, value in item.items():
                        paragraph = doc.add_paragraph(style="List Bullet")
                        paragraph.add_run(f"{self._format_label(key)}: ").bold = True
                        paragraph.add_run(str(value))
                else:
                    doc.add_paragraph(str(item), style="List Bullet")
            return

        doc.add_paragraph(str(content))

    def export_pdf(self, proposal: dict[str, Any], filename: str = "proposal.pdf") -> str:

        os.makedirs(settings.EXPORT_FOLDER, exist_ok=True)
        base_name, _ = os.path.splitext(filename)
        output_base = os.path.join(settings.EXPORT_FOLDER, base_name)
        output_path = f"{output_base}.pdf"

        latex_doc = LatexDocument(
            geometry_options={
                "margin": "1in"
            }
        )
        latex_doc.preamble.append(Command("title", "Project Proposal"))
        latex_doc.preamble.append(Command("date", Command("today")))
        latex_doc.append(Command("maketitle"))

        for section, content in proposal.items():
            with latex_doc.create(Section(self._escape(self._format_label(section)))) as section_block:
                self._write_latex_content(section_block, content, depth=1)

        try:
            latex_doc.generate_pdf(
                filepath=output_base,
                clean_tex=False,
                clean=True,
                compiler="pdflatex"
            )
        except CompilerError as error:
            raise RuntimeError(
                "PyLaTeX could not compile PDF. Install a LaTeX compiler (pdflatex/latexmk) and ensure it is in PATH."
            ) from error

        return output_path

    def _write_latex_content(self, container: LatexDocument | Section | Subsection, content: Any, depth: int) -> None:

        if isinstance(content, str):
            for paragraph in self._split_paragraphs(content):
                container.append(self._escape(paragraph))
                container.append("\n\n")
            return

        if isinstance(content, dict):
            for key, value in content.items():
                title = self._escape(key)

                if depth <= 1:
                    with container.create(Subsection(title)) as subsection_block:
                        self._write_latex_content(subsection_block, value, depth + 1)
                else:
                    container.append(Command("textbf", title))
                    container.append(": ")
                    self._write_latex_content(container, value, depth + 1)
            return

        if isinstance(content, list):
            with container.create(Itemize()) as itemize:
                for item in content:
                    if isinstance(item, dict):
                        joined = "; ".join(
                            f"{self._format_label(k)}: {v}" for k, v in item.items()
                        )
                        itemize.add_item(self._escape(joined))
                    else:
                        itemize.add_item(self._escape(str(item)))
            return

        container.append(self._escape(str(content)))
        container.append("\n")

    def export_pptx(self, proposal: dict[str, Any], filename: str = "proposal.pptx") -> str:

        path = os.path.join(settings.EXPORT_FOLDER, filename)
        prs = Presentation()

        for section, content in proposal.items():
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = slide.shapes.title
            body = slide.placeholders[1]

            title.text = self._format_label(section)
            body.text = str(content)

        prs.save(path)
        return path

    @staticmethod
    def _split_paragraphs(value: str) -> list[str]:
        return [chunk.strip() for chunk in value.split("\n\n") if chunk.strip()] or [value]

    @staticmethod
    def _format_label(label: Any) -> str:
        return str(label).replace("_", " ").title()

    @staticmethod
    def _escape(value: Any) -> str:
        return escape_latex(str(value))
